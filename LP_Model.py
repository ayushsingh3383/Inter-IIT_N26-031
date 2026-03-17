import os
import pulp
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

# -------------------------------
# 1. DATA SETUP
# -------------------------------
stations = np.arange(1, 16)
distance_step = 100  # meters between stations
width = 20      # pipeline clearing width in meters
area = distance_step * width  # 2000 m^2 per section

original_elevation = np.array([
    51, 67, 61, 55, 40, 45, 20, 22, 20, 42, 60, 42, 25, 28, 22
])

# Costs (in INR)
cost_cut = 150
cost_fill = 100
cost_dump = 70
cost_borrow = 120

# -------------------------------
# 2. INITIALIZE MODEL & VARIABLES
# -------------------------------
model = pulp.LpProblem("Pipeline_Leveling_Optimization", pulp.LpMinimize)

y = pulp.LpVariable.dicts("Elevation", stations, cat='Continuous')
cut = pulp.LpVariable.dicts("Cut", stations, lowBound=0, cat='Continuous')
fill = pulp.LpVariable.dicts("Fill", stations, lowBound=0, cat='Continuous')

borrow = pulp.LpVariable("Borrow", lowBound=0, cat='Continuous')
dump = pulp.LpVariable("Dump", lowBound=0, cat='Continuous')

# -------------------------------
# 3. OBJECTIVE FUNCTION
# -------------------------------
model += (
    cost_cut * area * pulp.lpSum(cut[i] for i in stations) +
    cost_fill * area * pulp.lpSum(fill[i] for i in stations) +
    cost_borrow * borrow +
    cost_dump * dump
), "Total_Cost"

# -------------------------------
# 4. CONSTRAINTS
# -------------------------------
for i in stations:
    g = original_elevation[i-1]
    model += y[i] == g - cut[i] + fill[i], f"Elev_Balance_St_{i}"

model += y[1] == 51, "Start_Boundary"
model += y[15] == 22, "End_Boundary"

for i in range(1, 15):
    model += y[i+1] - y[i] <= 5, f"Grade_Max_Pos_{i}"
    model += y[i+1] - y[i] >= -5, f"Grade_Max_Neg_{i}"

for i in range(2, 15):
    model += y[i+1] - 2*y[i] + y[i-1] <= 6, f"Grade_Change_Pos_{i}"
    model += y[i+1] - 2*y[i] + y[i-1] >= -6, f"Grade_Change_Neg_{i}"

model += y[2] - y[1] - 4 <= 6, "Init_Grade_Pos"
model += y[2] - y[1] - 4 >= -6, "Init_Grade_Neg"

model += -3 - (y[15] - y[14]) <= 6, "Final_Grade_Pos"
model += -3 - (y[15] - y[14]) >= -6, "Final_Grade_Neg"

model += (
    area * pulp.lpSum(cut[i] for i in stations) + borrow ==
    area * pulp.lpSum(fill[i] for i in stations) + dump
), "Material_Balance"

# -------------------------------
# 5. SOLVE & EXTRACT RESULTS
# -------------------------------
model.solve(pulp.PULP_CBC_CMD(msg=False))

opt_elev = [round(pulp.value(y[i]), 2) for i in stations]
cut_vals = [round(pulp.value(cut[i]), 2) for i in stations]
fill_vals = [round(pulp.value(fill[i]), 2) for i in stations]

# Calculate Extra Columns for the matching Excel layout
distances = [(i - 1) * 100 for i in stations]
station_costs = [(c * cost_cut + f * cost_fill) * area for c, f in zip(cut_vals, fill_vals)]

# Calculate Grades
grades = [4.0] # Initial incoming grade is 4%
for i in range(1, len(stations)):
    g_val = opt_elev[i] - opt_elev[i-1]
    grades.append(round(g_val, 2))

# Calculate Change in Grades
change_in_grades = [""] # Empty for the first station
for i in range(1, len(grades)):
    cg_val = grades[i] - grades[i-1]
    change_in_grades.append(round(cg_val, 2))

# Main Dataframe formatted EXACTLY like the image
df = pd.DataFrame({
    "Station Number": stations,
    "Distance along the pipeline (m)": distances,
    "Original Elevation (m)": original_elevation,
    "Optimized Elevation (m)": opt_elev,
    "Cut (m)": cut_vals,
    "Fill (m)": fill_vals,
    "Cost": station_costs,
    "Grade (%)": grades,
    "Change in Grade (%)": change_in_grades
})

# Summary Dataframe
summary_df = pd.DataFrame({
    "Project Metric": [
        "Total Volume Cut (m³)",
        "Total Volume Filled (m³)",
        "Extra Volume Borrowed (m³)",
        "Excess Volume Dumped (m³)",
        "LOWEST TOTAL PROJECT COST (₹)"
    ],
    "Value": [
        sum(cut_vals) * area,
        sum(fill_vals) * area,
        pulp.value(borrow),
        pulp.value(dump),
        pulp.value(model.objective)
    ]
})

# -------------------------------
# 6. EXPORT TO EXCEL WITH STYLING
# -------------------------------
output_dir = r"C:\Users\Ayushi Singh\LP_solver_hydraulics"
os.makedirs(output_dir, exist_ok=True) 
excel_filepath = os.path.join(output_dir, "Optimized_Vertical_Profile.xlsx")

try:
    with pd.ExcelWriter(excel_filepath, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='Optimized Profile', index=False)
        summary_df.to_excel(writer, sheet_name='Project Summary', index=False)
        
        # Get the workbook and the main worksheet
        workbook = writer.book
        worksheet = writer.sheets['Optimized Profile']
        
        # Define styles
        header_fill = PatternFill(start_color="92D050", end_color="92D050", fill_type="solid") # Green
        cut_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")    # Light Peach/Red
        fill_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")   # Light Blue
        cost_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")   # Light Yellow
        
        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), 
                             top=Side(style='thin'), bottom=Side(style='thin'))
        center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
        # 1. Apply styles to the Header Row
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = Font(bold=True)
            cell.alignment = center_align
            cell.border = thin_border
            
        # 2. Apply styles and conditional formatting to the Data Rows
        for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row, min_col=1, max_col=worksheet.max_column):
            for cell in row:
                cell.alignment = center_align
                cell.border = thin_border
                
                # Get the column name to apply conditional color highlighting
                col_name = worksheet.cell(row=1, column=cell.column).value
                
                if col_name == "Cut (m)" and isinstance(cell.value, (int, float)) and cell.value > 0:
                    cell.fill = cut_fill
                elif col_name == "Fill (m)" and isinstance(cell.value, (int, float)) and cell.value > 0:
                    cell.fill = fill_fill
                elif col_name == "Cost" and isinstance(cell.value, (int, float)) and cell.value > 0:
                    cell.fill = cost_fill

        # 3. Adjust Column Widths for better readability
        for col in worksheet.columns:
            col_letter = col[0].column_letter
            if col_letter == 'B': # Distance column needs more space
                worksheet.column_dimensions[col_letter].width = 30
            elif col_letter in ['H', 'I']: # Grade columns
                worksheet.column_dimensions[col_letter].width = 20
            else:
                worksheet.column_dimensions[col_letter].width = 18

    print("-" * 50)
    print(f"Optimization Status: {pulp.LpStatus[model.status]}")
    print(f"LOWEST TOTAL COST: ₹ {pulp.value(model.objective):,.2f}")
    print("-" * 50)
    print(f"[+] Successfully saved exactly styled Excel file to:\n    {excel_filepath}")

except PermissionError:
    print("-" * 50)
    print(f"\n[!] ERROR: Could not save the file because it is currently open in Excel.")
    print(f"    Please close '{excel_filepath}' and run the script again!\n")
    print("-" * 50)

# -------------------------------
# 7. VISUALIZATION
# -------------------------------
plt.figure(figsize=(12, 6))
plt.plot(distances, original_elevation, marker='o', linestyle='--', color='gray', label='Original Terrain')
plt.plot(distances, opt_elev, marker='s', linestyle='-', color='blue', linewidth=2, label='Optimized Pipeline Profile')
plt.fill_between(distances, original_elevation, opt_elev, where=(np.array(opt_elev) > original_elevation), interpolate=True, color='green', alpha=0.3, label='Fill Area')
plt.fill_between(distances, original_elevation, opt_elev, where=(np.array(opt_elev) < original_elevation), interpolate=True, color='red', alpha=0.3, label='Cut Area')
plt.title('Pipeline Vertical Profile Optimization', fontsize=16, fontweight='bold')
plt.xlabel('Distance along pipeline (m)', fontsize=12)
plt.ylabel('Elevation (m)', fontsize=12)
plt.xticks(distances, labels=[f"{d}" for d in distances], rotation=45)
plt.grid(True, linestyle=':', alpha=0.7)
plt.legend(loc='best', fontsize=11)
plt.tight_layout()
plt.show(block=False)
plt.pause(0.1)
input("Press Enter in the console to close the graph and exit the script...")
